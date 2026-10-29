import argparse
import io
import json
import os
import re
import tempfile
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import quote, urlparse

import requests
from openpyxl import load_workbook

MAX_BYTES = 25 * 1024 * 1024
DATASETS = {'alerts', 'projects', 'qualityGates', 'resources', 'inspections', 'commercial'}


def utc_now():
    return datetime.now(timezone.utc).isoformat()


def acquire_token():
    import msal
    tenant = os.environ.get('M365_TENANT_ID', '')
    client = os.environ.get('M365_CLIENT_ID', '')
    if not re.fullmatch(r'[0-9a-fA-F-]{36}', tenant) or not client:
        raise ValueError('Configure M365_TENANT_ID and M365_CLIENT_ID')
    certificate = os.environ.get('M365_CERTIFICATE_PATH')
    if certificate:
        credential = {
            'private_key': Path(certificate).read_text(),
            'thumbprint': os.environ['M365_CERTIFICATE_THUMBPRINT'],
        }
    else:
        credential = os.environ.get('M365_CLIENT_SECRET')
        if not credential:
            raise ValueError('Configure a certificate or M365_CLIENT_SECRET')
    app = msal.ConfidentialClientApplication(
        client, authority=f'https://login.microsoftonline.com/{tenant}',
        client_credential=credential,
    )
    result = app.acquire_token_for_client(scopes=['https://graph.microsoft.com/.default'])
    if 'access_token' not in result:
        raise RuntimeError('Microsoft authentication failed; check consent and credentials')
    return result['access_token']


def download(source, token):
    drive = quote(source['driveId'], safe='')
    item = quote(source['itemId'], safe='')
    url = f'https://graph.microsoft.com/v1.0/drives/{drive}/items/{item}/content'
    response = requests.get(url, headers={'Authorization': f'Bearer {token}'},
                            timeout=(10, 60), allow_redirects=False, stream=True)
    if response.status_code in (301, 302, 303, 307, 308):
        location = response.headers.get('Location', '')
        response.close()
        parsed = urlparse(location)
        host = parsed.hostname or ''
        if parsed.scheme != 'https' or parsed.username or parsed.password or not (
            host.endswith('.sharepoint.com') or host.endswith('.1drv.com')
            or host == 'graph.microsoft.com'
        ):
            raise ValueError('Download host is not allowed; review provider configuration')
        response = requests.get(location, timeout=(10, 60), allow_redirects=False, stream=True)
    with response:
        if response.status_code != 200:
            raise RuntimeError(f'Workbook download failed: HTTP {response.status_code}')
        chunks, size = [], 0
        for chunk in response.iter_content(65536):
            size += len(chunk)
            if size > MAX_BYTES:
                raise ValueError('Workbook exceeds the 25 MiB limit')
            chunks.append(chunk)
        return b''.join(chunks)


def serialise(value, field):
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    if field == 'Due' and value not in (None, ''):
        text = str(value).strip()
        for pattern in ('%d/%m/%Y', '%Y-%m-%d'):
            try:
                return datetime.strptime(text, pattern).date().isoformat()
            except ValueError:
                pass
        raise ValueError('Due dates must be Excel dates, DD/MM/YYYY or YYYY-MM-DD')
    return value


def parse_workbook(content, source):
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    output, seen = [], set()
    try:
        for mapping in source['sheets']:
            dataset = mapping['dataset']
            if dataset not in DATASETS:
                raise ValueError('Unsupported dataset')
            sheet_name = mapping['name']
            if sheet_name not in workbook.sheetnames:
                raise ValueError('Configured worksheet is missing')
            columns = mapping['columns']
            id_field = mapping['idField']
            if id_field not in columns:
                raise ValueError('ID field must be mapped')
            rows = workbook[sheet_name].iter_rows(values_only=True)
            headers = [str(v).strip() if v is not None else '' for v in next(rows, ())]
            for column in columns.values():
                if headers.count(column) != 1:
                    raise ValueError('Required column is missing or duplicated')
            positions = {field: headers.index(column) for field, column in columns.items()}
            for row_number, row in enumerate(rows, 2):
                if not any(v is not None and v != '' for v in row):
                    continue
                record = {field: serialise(row[index] if index < len(row) else None, field)
                          for field, index in positions.items()}
                identifier = record.get(id_field)
                if identifier is None or str(identifier).strip() == '':
                    raise ValueError(f'Missing record ID at row {row_number}')
                identity = json.dumps([source['id'], sheet_name, str(identifier)], ensure_ascii=False)
                if identity in seen:
                    raise ValueError('Duplicate record ID within worksheet')
                seen.add(identity)
                output.append({'id': identity, 'dataset': dataset, 'sourceId': source['id'],
                               'sheet': sheet_name, 'data': record})
    finally:
        workbook.close()
    return output


def save_atomic(path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(dir=path.parent, prefix='.sync-')
    try:
        with os.fdopen(fd, 'w') as handle:
            json.dump(payload, handle, ensure_ascii=False, default=str)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def run(config_path, state_path):
    config = json.loads(Path(config_path).read_text())
    sources = [s for s in config['sources'] if s.get('enabled', False)]
    ids = [s['id'] for s in sources]
    if len(set(ids)) != len(ids):
        raise ValueError('Source IDs must be unique')
    path = Path(state_path)
    previous = json.loads(path.read_text()) if path.exists() else {'sources': {}}
    state = {'updatedAt': utc_now(), 'sources': {}}
    token = acquire_token() if sources else None
    failures = 0
    for source in sources:
        old = previous.get('sources', {}).get(source['id'], {})
        try:
            records = parse_workbook(download(source, token), source)
            state['sources'][source['id']] = {
                'status': 'ok', 'lastAttempt': utc_now(), 'lastSuccess': utc_now(),
                'records': records,
            }
        except Exception as error:
            failures += 1
            state['sources'][source['id']] = {
                'status': 'error', 'lastAttempt': utc_now(),
                'lastSuccess': old.get('lastSuccess'), 'records': old.get('records', []),
                'error': type(error).__name__,
            }
    save_atomic(path, state)
    print(json.dumps({'enabledSources': len(sources), 'failedSources': failures}))
    return 1 if failures else 0


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Private Microsoft 365 workbook sync')
    parser.add_argument('--config', default='backend/sources.local.json')
    parser.add_argument('--state', default='.private-sync/state.json')
    args = parser.parse_args()
    raise SystemExit(run(args.config, args.state))
